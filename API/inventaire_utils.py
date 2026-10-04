"""
Outils de l'inventaire : liste de comptage (feuille imprimable et saisie en masse),
saisie en masse des quantites comptees et export Excel d'une session.

La feuille de comptage et l'ecran de saisie en masse utilisent la MEME liste ordonnee
(`construire_liste_comptage`) : les numeros de ligne imprimes correspondent donc aux
numeros affiches a l'ecran, ce qui permet de ressaisir la feuille papier sans chercher.
Aucune quantite theorique n'y figure (comptage a l'aveugle).
"""
from io import BytesIO

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from .models import Categorie, InventoryLine, Produit


def produits_du_tenant(request):
    """Produits actifs visibles par l'utilisateur (meme regle que TenantFilterMixin)."""
    qs = Produit.objects.filter(is_active=True)
    user = request.user
    if user.is_superuser or user.is_staff:
        return qs
    company = getattr(request, 'company', None)
    if company is not None:
        return qs.filter(Q(company=company) | Q(company__isnull=True))
    return qs.filter(company__isnull=True)


def _ids_categorie(categorie_id):
    """Ids de la categorie et de toutes ses sous-categories ([] si introuvable)."""
    try:
        categorie = Categorie.objects.get(pk=int(categorie_id))
    except (Categorie.DoesNotExist, TypeError, ValueError):
        return []
    return [categorie.id] + [c.id for c in categorie.get_all_children()]


def construire_liste_comptage(session, request, categorie_id=None, restants=False):
    """
    Liste ordonnee des produits a compter : tous les produits actifs du tenant, plus les
    produits deja presents dans la session (meme desactives depuis). Tri par categorie
    (chemin complet) puis designation. Chaque element porte un `numero` sequentiel.
    """
    lignes = {l.produit_id: l for l in session.lignes.select_related('counted_by')}

    produits = produits_du_tenant(request)
    if lignes:
        produits = Produit.objects.filter(Q(pk__in=produits.values('pk')) | Q(pk__in=list(lignes.keys())))
    if categorie_id:
        produits = produits.filter(categorie_id__in=_ids_categorie(categorie_id))
    produits = produits.select_related('categorie', 'categorie__parent').prefetch_related('conditionnements')

    elements = []
    for p in produits:
        ligne = lignes.get(p.id)
        compte = ligne.counted_qty if ligne is not None else None
        if restants and compte is not None:
            continue
        categorie = p.categorie.get_full_path() if p.categorie_id else 'Sans catégorie'
        elements.append({
            'produit_id': p.id,
            'reference': p.reference or '',
            'code_barre': p.code_barre or '',
            'designation': p.designation or '',
            'categorie': categorie,
            'unite_mesure': p.get_unite_mesure_display() if hasattr(p, 'get_unite_mesure_display') else '',
            'unites_par_carton': p.get_unites_par_carton(),
            'line_id': ligne.id if ligne is not None else None,
            'counted_qty': compte,
            'counted_by_username': ligne.counted_by.username if ligne is not None and ligne.counted_by_id else None,
        })
    elements.sort(key=lambda e: (e['categorie'].lower(), e['designation'].lower(), e['reference'].lower()))
    for i, e in enumerate(elements, start=1):
        e['numero'] = i
    return elements


def grouper_par_categorie(elements):
    """[(categorie, [elements])] en conservant l'ordre."""
    groupes = []
    for e in elements:
        if not groupes or groupes[-1][0] != e['categorie']:
            groupes.append((e['categorie'], []))
        groupes[-1][1].append(e)
    return groupes


def _trouver_produit(entree, produits_qs, cache):
    """Resout le produit d'une ligne de saisie par id, reference ou code-barres."""
    pid = entree.get('produit_id') or entree.get('produit')
    if pid not in (None, ''):
        try:
            return produits_qs.get(pk=int(pid)), None
        except (Produit.DoesNotExist, TypeError, ValueError):
            return None, f"Produit #{pid} introuvable"
    cle = str(entree.get('reference') or entree.get('code_barre') or '').strip()
    if not cle:
        return None, "Produit non précisé (produit_id, reference ou code_barre)"
    if cle.lower() in cache:
        return cache[cle.lower()], None
    trouves = list(produits_qs.filter(Q(reference__iexact=cle) | Q(code_barre__iexact=cle))[:2])
    if not trouves:
        return None, f"Aucun produit avec la référence ou le code-barres « {cle} »"
    if len(trouves) > 1:
        return None, f"Plusieurs produits correspondent à « {cle} »"
    cache[cle.lower()] = trouves[0]
    return trouves[0], None


def saisir_en_masse(session, entrees, request):
    """
    Enregistre en une fois des quantites comptees.
    `entrees` : liste de {produit_id | reference | code_barre, counted_qty}.
    counted_qty = null remet la ligne a « non compte ».
    Les lignes valides sont enregistrees ; les lignes en erreur sont listees dans le
    rapport, avec leur position (1 = premiere ligne envoyee).
    """
    rapport = {'creees': 0, 'mises_a_jour': 0, 'reinitialisees': 0, 'inchangees': 0, 'erreurs': []}
    if not isinstance(entrees, list):
        rapport['erreurs'].append({'position': 0, 'message': "Le champ « lignes » doit être une liste."})
        return rapport

    # Produits autorises : ceux du tenant + ceux deja dans la session
    autorises = Produit.objects.filter(
        Q(pk__in=produits_du_tenant(request).values('pk')) | Q(pk__in=session.lignes.values('produit_id')))
    cache = {}
    maintenant = timezone.now()
    vus = set()

    with transaction.atomic():
        for position, entree in enumerate(entrees, start=1):
            if not isinstance(entree, dict):
                rapport['erreurs'].append({'position': position, 'message': 'Ligne invalide'})
                continue
            produit, erreur = _trouver_produit(entree, autorises, cache)
            etiquette = str(entree.get('reference') or entree.get('code_barre') or entree.get('produit_id') or '')
            if erreur:
                rapport['erreurs'].append({'position': position, 'reference': etiquette, 'message': erreur})
                continue
            if produit.id in vus:
                rapport['erreurs'].append({'position': position, 'reference': produit.reference,
                                           'message': f"« {produit.reference} » apparaît plusieurs fois dans la saisie"})
                continue

            brut = entree.get('counted_qty', entree.get('quantite'))
            if brut in (None, ''):
                quantite = None
            else:
                try:
                    valeur = float(str(brut).strip().replace(' ', '').replace(',', '.'))
                    if valeur != int(valeur):
                        raise ValueError
                    quantite = int(valeur)
                except (TypeError, ValueError, OverflowError):
                    rapport['erreurs'].append({'position': position, 'reference': produit.reference,
                                               'message': f"Quantité invalide « {brut} » (nombre entier attendu)"})
                    continue
                if quantite < 0:
                    rapport['erreurs'].append({'position': position, 'reference': produit.reference,
                                               'message': 'La quantité ne peut pas être négative'})
                    continue
            vus.add(produit.id)

            ligne = InventoryLine.objects.filter(session=session, produit=produit).first()
            if ligne is None:
                if quantite is None:
                    rapport['inchangees'] += 1
                    continue
                InventoryLine.objects.create(session=session, produit=produit, snapshot_qty=produit.quantite or 0,
                                             counted_qty=quantite, counted_by=request.user, counted_at=maintenant)
                rapport['creees'] += 1
            elif ligne.counted_qty == quantite:
                rapport['inchangees'] += 1
            elif quantite is None:
                ligne.counted_qty = None
                ligne.counted_by = None
                ligne.counted_at = None
                ligne.save()
                rapport['reinitialisees'] += 1
            else:
                ligne.counted_qty = quantite
                ligne.counted_by = request.user
                ligne.counted_at = maintenant
                ligne.save()
                rapport['mises_a_jour'] += 1

        session.update_completion_percentage()
        if session.statut == 'draft' and (rapport['creees'] or rapport['mises_a_jour']):
            session.statut = 'in_progress'
            session.save(update_fields=['statut'])
    return rapport


def _elements_export(session, request):
    """Lignes a exporter. Tant que la session est ouverte (draft / in_progress), on liste TOUS
    les produits a compter (comme la feuille de comptage) afin que l'agent puisse remplir la
    quantite reelle des produits pas encore comptes ; les lignes deja saisies sont fusionnees.
    Une session cloturee n'exporte que ses lignes (etat fige). Sans `request`, on ne peut pas
    resoudre les produits du tenant : on retombe sur les lignes de la session."""
    lignes = {l.produit_id: l for l in session.lignes.select_related(
        'produit', 'produit__categorie', 'produit__categorie__parent', 'counted_by')}
    ouverte = session.statut in ('draft', 'in_progress')
    if ouverte and request is not None:
        produits = produits_du_tenant(request)
        if lignes:
            produits = Produit.objects.filter(Q(pk__in=produits.values('pk')) | Q(pk__in=list(lignes.keys())))
        produits = list(produits.select_related('categorie', 'categorie__parent'))
    else:
        produits = [l.produit for l in lignes.values()]

    elements = []
    for p in produits:
        l = lignes.get(p.id)
        elements.append({
            'produit': p,
            'categorie': p.categorie.get_full_path() if p.categorie_id else '',
            'theorique': l.snapshot_qty if l is not None else (p.quantite or 0),
            'compte': l.counted_qty if l is not None else None,
            'ecart': l.get_variance() if l is not None else None,
            'compte_par': l.counted_by.username if l is not None and l.counted_by_id else '',
            'compte_le': timezone.localtime(l.counted_at).strftime('%d/%m/%Y %H:%M') if l is not None and l.counted_at else '',
        })
    elements.sort(key=lambda e: (e['categorie'].lower(), (e['produit'].designation or '').lower(),
                                 (e['produit'].reference or '').lower()))
    return elements


def exporter_session_excel(session, request=None):
    """Classeur Excel d'une session : en-tete, lignes avec theorique / compte / ecart.
    Voir `_elements_export` pour le choix des produits listes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Inventaire'

    gras = Font(bold=True)
    entete_fill = PatternFill('solid', fgColor='2D3748')
    bordure = Border(*(Side(style='thin', color='CBD5E0'),) * 4)

    ws['A1'] = f"Inventaire {session.numero}"
    ws['A1'].font = Font(bold=True, size=14)
    infos = [
        ('Date', session.date.strftime('%d/%m/%Y') if session.date else ''),
        ('Statut', session.get_statut_display()),
        ('Créé par', session.created_by.username if session.created_by_id else ''),
        ('Validé par', session.validated_by.username if session.validated_by_id else ''),
        ('Avancement', f"{session.completed_products}/{session.total_products} ({float(session.completion_percentage or 0):.0f} %)"),
        ('Note', session.note or ''),
    ]
    for i, (libelle, valeur) in enumerate(infos, start=2):
        ws.cell(row=i, column=1, value=libelle).font = gras
        ws.cell(row=i, column=2, value=valeur)

    colonnes = ['N°', 'Référence', 'Code-barres', 'Désignation', 'Catégorie', 'Stock théorique',
                'Quantité comptée', 'Écart', 'Compté par', 'Date comptage']
    ligne_entete = len(infos) + 3
    for c, titre in enumerate(colonnes, start=1):
        cell = ws.cell(row=ligne_entete, column=c, value=titre)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = entete_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = bordure

    rouge, vert = Font(color='C53030', bold=True), Font(color='2F855A', bold=True)
    a_saisir = PatternFill('solid', fgColor='FEFCBF')   # jaune : quantite reelle a remplir
    r = ligne_entete
    total_ecart = 0
    for n, e in enumerate(_elements_export(session, request), start=1):
        r += 1
        p, ecart = e['produit'], e['ecart']
        if ecart:
            total_ecart += ecart
        valeurs = [
            n, p.reference, p.code_barre, p.designation, e['categorie'],
            e['theorique'], e['compte'], ecart, e['compte_par'], e['compte_le'],
        ]
        for c, v in enumerate(valeurs, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = bordure
            if c == 7 and e['compte'] is None:
                cell.fill = a_saisir
            if c == 8 and ecart:
                cell.font = vert if ecart > 0 else rouge
    r += 1
    ws.cell(row=r, column=5, value='Total des écarts').font = gras
    ws.cell(row=r, column=8, value=total_ecart).font = gras

    for c, largeur in enumerate([6, 16, 18, 40, 28, 16, 16, 10, 16, 18], start=1):
        ws.column_dimensions[get_column_letter(c)].width = largeur
    ws.freeze_panes = ws.cell(row=ligne_entete + 1, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

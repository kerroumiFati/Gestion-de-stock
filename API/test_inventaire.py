"""Tests de l'inventaire : validation incomplete (ex erreur 500), feuille de comptage sans
quantite theorique, saisie en masse et export Excel de session."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from rest_framework.test import APIClient

from .models import Categorie, Conditionnement, InventoryLine, InventorySession, Produit, StockMove


class InventaireTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('admin', password='x', is_staff=True)
        self.api = APIClient()
        self.api.force_authenticate(self.user)
        self.boissons = Categorie.objects.create(nom='Boissons')
        self.jus = Categorie.objects.create(nom='Jus', parent=self.boissons)
        self.snacks = Categorie.objects.create(nom='Snacks')
        # Quantites theoriques volontairement reconnaissables (ne doivent pas figurer sur la feuille)
        self.eau = Produit.objects.create(reference='EAU-01', code_barre='3700000000011', designation='Eau 1L',
                                          prixU=Decimal('10'), categorie=self.boissons, quantite=98765)
        self.orange = Produit.objects.create(reference='JUS-01', code_barre='3700000000028', designation='Jus orange',
                                             prixU=Decimal('20'), categorie=self.jus, quantite=87654)
        self.chips = Produit.objects.create(reference='CHIPS-01', code_barre='3700000000035', designation='Chips sel',
                                            prixU=Decimal('5'), categorie=self.snacks, quantite=76543)
        Conditionnement.objects.create(produit=self.eau, unites_par_carton=12)
        self.session = InventorySession.objects.create(numero='INV-2026-0001', created_by=self.user)
        self.url = f'/API/inventaires/{self.session.id}/'

    # ------------------------------------------------------------ validation
    def test_validation_sans_comptage_renvoie_400_et_non_500(self):
        r = self.api.post(self.url + 'validate/')
        self.assertEqual(r.status_code, 400, r.content)
        self.assertIn('Aucun produit', r.json()['detail'])
        self.assertEqual(r.json()['missing_products'], 0)

    def test_validation_incomplete_liste_les_produits_manquants(self):
        InventoryLine.objects.create(session=self.session, produit=self.eau, snapshot_qty=100, counted_qty=90)
        InventoryLine.objects.create(session=self.session, produit=self.chips, snapshot_qty=50, counted_qty=None)
        r = self.api.post(self.url + 'validate/')
        self.assertEqual(r.status_code, 400, r.content)
        data = r.json()
        self.assertEqual(data['missing_products'], 1)
        self.assertEqual(data['missing_products_list'][0]['reference'], 'CHIPS-01')
        self.assertEqual(data['completion_percentage'], 50.0)
        self.session.refresh_from_db()
        self.assertNotEqual(self.session.statut, 'validated')
        self.assertEqual(StockMove.objects.count(), 0)   # aucun ajustement partiel

    def test_validation_complete_ajuste_le_stock(self):
        InventoryLine.objects.create(session=self.session, produit=self.eau, snapshot_qty=98765, counted_qty=98760)
        r = self.api.post(self.url + 'validate/')
        self.assertEqual(r.status_code, 200, r.content)
        self.eau.refresh_from_db()
        self.assertEqual(self.eau.quantite, 98760)
        self.session.refresh_from_db()
        self.assertEqual(self.session.statut, 'validated')
        # une session validee ne se revalide pas
        self.assertEqual(self.api.post(self.url + 'validate/').status_code, 400)

    # ------------------------------------------------------------ liste et feuille de comptage
    def test_liste_comptage_ordonnee_numerotee_sans_theorique(self):
        r = self.api.get(self.url + 'liste_comptage/')
        self.assertEqual(r.status_code, 200, r.content)
        lignes = r.json()['lignes']
        self.assertEqual([l['reference'] for l in lignes], ['EAU-01', 'JUS-01', 'CHIPS-01'])
        self.assertEqual([l['numero'] for l in lignes], [1, 2, 3])
        self.assertEqual(lignes[0]['unites_par_carton'], 12)
        self.assertEqual(lignes[1]['categorie'], 'Boissons > Jus')
        for l in lignes:
            self.assertNotIn('snapshot_qty', l)
            self.assertNotIn('quantite', l)

    def test_liste_comptage_filtres(self):
        InventoryLine.objects.create(session=self.session, produit=self.eau, snapshot_qty=1, counted_qty=5)
        r = self.api.get(self.url + f'liste_comptage/?categorie={self.boissons.id}')
        self.assertEqual([l['reference'] for l in r.json()['lignes']], ['EAU-01', 'JUS-01'])   # sous-categorie incluse
        r = self.api.get(self.url + 'liste_comptage/?restants=1')
        self.assertEqual([l['reference'] for l in r.json()['lignes']], ['JUS-01', 'CHIPS-01'])

    def test_feuille_comptage_imprimable_sans_quantite_theorique(self):
        InventoryLine.objects.create(session=self.session, produit=self.chips, snapshot_qty=76543, counted_qty=None)
        r = self.api.get(self.url + 'feuille_comptage/')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r['Content-Type'].startswith('text/html'))
        html = r.content.decode('utf-8')
        self.assertIn('FEUILLE DE COMPTAGE', html)
        self.assertIn('INV-2026-0001', html)
        for attendu in ('EAU-01', 'Jus orange', '3700000000035', 'ct de 12', 'window.print()'):
            self.assertIn(attendu, html)
        for interdit in ('98765', '87654', '76543', 'héorique'):
            self.assertNotIn(interdit, html)
        # filtre "restants"
        InventoryLine.objects.create(session=self.session, produit=self.eau, snapshot_qty=1, counted_qty=3)
        html = self.api.get(self.url + 'feuille_comptage/?restants=1').content.decode('utf-8')
        self.assertNotIn('EAU-01', html)
        self.assertIn('CHIPS-01', html)

    # ------------------------------------------------------------ saisie en masse
    def test_saisie_masse_cree_modifie_et_efface(self):
        InventoryLine.objects.create(session=self.session, produit=self.chips, snapshot_qty=50, counted_qty=40)
        r = self.api.post(self.url + 'saisie_masse/', {'lignes': [
            {'produit_id': self.eau.id, 'counted_qty': 120},        # creation
            {'reference': 'jus-01', 'counted_qty': '7'},            # par reference, insensible a la casse
            {'code_barre': '3700000000035', 'counted_qty': 45},     # mise a jour par code-barres
        ]}, format='json')
        self.assertEqual(r.status_code, 200, r.content)
        rapport = r.json()['rapport']
        self.assertEqual((rapport['creees'], rapport['mises_a_jour'], rapport['erreurs']), (2, 1, []))
        self.assertEqual(InventoryLine.objects.get(session=self.session, produit=self.eau).counted_qty, 120)
        self.assertEqual(InventoryLine.objects.get(session=self.session, produit=self.eau).snapshot_qty, 98765)
        self.assertEqual(InventoryLine.objects.get(session=self.session, produit=self.orange).counted_qty, 7)
        ligne = InventoryLine.objects.get(session=self.session, produit=self.chips)
        self.assertEqual((ligne.counted_qty, ligne.snapshot_qty, ligne.counted_by), (45, 50, self.user))
        self.session.refresh_from_db()
        self.assertEqual(self.session.statut, 'in_progress')
        self.assertEqual(self.session.completed_products, 3)

        # ressaisie : correction + remise a « non compte » + ligne inchangee
        r = self.api.post(self.url + 'saisie_masse/', {'lignes': [
            {'produit_id': self.eau.id, 'counted_qty': 118},
            {'produit_id': self.orange.id, 'counted_qty': None},
            {'produit_id': self.chips.id, 'counted_qty': 45},
        ]}, format='json')
        rapport = r.json()['rapport']
        self.assertEqual((rapport['mises_a_jour'], rapport['reinitialisees'], rapport['inchangees']), (1, 1, 1))
        self.assertIsNone(InventoryLine.objects.get(session=self.session, produit=self.orange).counted_qty)
        self.assertEqual(r.json()['session']['completed_products'], 2)

    def test_saisie_masse_erreurs_signalees_lignes_valides_enregistrees(self):
        r = self.api.post(self.url + 'saisie_masse/', {'lignes': [
            {'reference': 'EAU-01', 'counted_qty': 10},
            {'reference': 'INCONNU', 'counted_qty': 5},
            {'reference': 'JUS-01', 'counted_qty': -3},
            {'reference': 'CHIPS-01', 'counted_qty': 'abc'},
            {'reference': 'EAU-01', 'counted_qty': 11},     # doublon
            {'reference': 'CHIPS-01', 'counted_qty': '2,5'},  # non entier
        ]}, format='json')
        self.assertEqual(r.status_code, 200, r.content)
        rapport = r.json()['rapport']
        self.assertEqual(rapport['creees'], 1)
        self.assertEqual([e['position'] for e in rapport['erreurs']], [2, 3, 4, 5, 6])
        self.assertEqual(InventoryLine.objects.filter(session=self.session).count(), 1)
        self.assertEqual(InventoryLine.objects.get(session=self.session).counted_qty, 10)

    def test_saisie_masse_refusee_si_vide_ou_session_validee(self):
        self.assertEqual(self.api.post(self.url + 'saisie_masse/', {'lignes': []}, format='json').status_code, 400)
        self.assertEqual(self.api.post(self.url + 'saisie_masse/', {}, format='json').status_code, 400)
        self.session.statut = 'validated'
        self.session.save()
        r = self.api.post(self.url + 'saisie_masse/', {'lignes': [{'produit_id': self.eau.id, 'counted_qty': 1}]}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(InventoryLine.objects.count(), 0)

    # ------------------------------------------------------------ export
    def test_export_excel_de_la_session(self):
        from openpyxl import load_workbook
        InventoryLine.objects.create(session=self.session, produit=self.eau, snapshot_qty=100, counted_qty=90,
                                     counted_by=self.user)
        InventoryLine.objects.create(session=self.session, produit=self.chips, snapshot_qty=50, counted_qty=55)
        r = self.api.get(self.url + 'export/')
        self.assertEqual(r.status_code, 200, r.content[:300])
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('inventaire_INV-2026-0001.xlsx', r['Content-Disposition'])
        ws = load_workbook(BytesIO(r.content)).active
        cellules = [[c.value for c in row] for row in ws.iter_rows()]
        self.assertEqual(cellules[0][0], 'Inventaire INV-2026-0001')
        entete = next(i for i, row in enumerate(cellules) if row[0] == 'N°')
        self.assertEqual(cellules[entete][5:8], ['Stock théorique', 'Quantité comptée', 'Écart'])
        self.assertEqual(cellules[entete + 1][1:8], ['EAU-01', '3700000000011', 'Eau 1L', 'Boissons', 100, 90, -10])
        # Session ouverte : le jus, pas encore compte, figure quand meme (theorique = stock actuel, compte vide)
        self.assertEqual(cellules[entete + 2][1], 'JUS-01')
        self.assertEqual(cellules[entete + 2][5:8], [87654, None, None])
        self.assertEqual(cellules[entete + 3][1], 'CHIPS-01')
        self.assertEqual(cellules[entete + 3][7], 5)
        self.assertEqual(cellules[entete + 4][7], -5)   # total des ecarts (lignes non comptees ignorees)

    def test_export_excel_session_vierge_liste_tous_les_produits(self):
        """Une session sans aucune ligne doit quand meme exporter tous les produits a compter
        (l'agent remplit la colonne « Quantite comptee » dans le fichier)."""
        from openpyxl import load_workbook
        r = self.api.get(self.url + 'export/')
        self.assertEqual(r.status_code, 200, r.content[:300])
        ws = load_workbook(BytesIO(r.content)).active
        cellules = [[c.value for c in row] for row in ws.iter_rows()]
        entete = next(i for i, row in enumerate(cellules) if row[0] == 'N°')
        references = [row[1] for row in cellules[entete + 1:entete + 4]]
        self.assertEqual(references, ['EAU-01', 'JUS-01', 'CHIPS-01'])
        self.assertEqual([row[6] for row in cellules[entete + 1:entete + 4]], [None, None, None])
        self.assertEqual(cellules[entete + 1][5], 98765)   # stock theorique = stock actuel du produit
        self.assertEqual(cellules[entete + 4][4], 'Total des écarts')

    def test_export_excel_session_validee_ne_liste_que_ses_lignes(self):
        from openpyxl import load_workbook
        InventoryLine.objects.create(session=self.session, produit=self.eau, snapshot_qty=100, counted_qty=90)
        self.session.statut = 'validated'
        self.session.save()
        r = self.api.get(self.url + 'export/')
        ws = load_workbook(BytesIO(r.content)).active
        cellules = [[c.value for c in row] for row in ws.iter_rows()]
        entete = next(i for i, row in enumerate(cellules) if row[0] == 'N°')
        self.assertEqual(cellules[entete + 1][1], 'EAU-01')
        self.assertEqual(cellules[entete + 2][4], 'Total des écarts')

    def test_details_de_session_exposent_la_categorie(self):
        InventoryLine.objects.create(session=self.session, produit=self.orange, snapshot_qty=1, counted_qty=0)
        ligne = self.api.get(self.url).json()['lignes'][0]
        self.assertEqual(ligne['produit_categorie'], 'Jus')
        self.assertTrue(ligne['is_completed'])   # 0 compte = compte

    def test_acces_anonyme_refuse(self):
        anonyme = APIClient()
        for suffixe in ('feuille_comptage/', 'liste_comptage/', 'export/'):
            self.assertIn(anonyme.get(self.url + suffixe).status_code, (401, 403))
